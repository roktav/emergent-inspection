#!/usr/bin/env python3
"""Build the ITI-branded Indonesian kickoff deck (16:9).

Rebuild:
    python3 docs/panduan-pengguna/build_intro_pptx.py

Visual tokens match print.css: ink #231915, accent #98A7D5, tint #D6D2E9,
paper white. Headings Poppins SemiBold, body Roboto — family names are set
even if the presenting PC lacks the files (same approach as the PDF).
"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
IMAGES = ROOT / "images"
LOGO = IMAGES / "brand" / "logo.png"
OUT = ROOT.parent / "Inspeksi_DT_Pengenalan.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(0x23, 0x19, 0x15)
ACCENT = RGBColor(0x98, 0xA7, 0xD5)
TINT = RGBColor(0xD6, 0xD2, 0xE9)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5C, 0x56, 0x62)
WARN = RGBColor(0x8A, 0x3B, 0x2A)

FONT_HEAD = "Poppins"
FONT_BODY = "Roboto"

MARGIN_X = Inches(0.55)
FOOTER_Y = Inches(7.12)


def set_run_font(run, name: str, size_pt: float, *, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def no_line(shape):
    shape.line.fill.background()


def set_slide_bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill: RGBColor, *, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    no_line(shape)
    if rounded:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def add_oval(slide, left, top, size, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    no_line(shape)
    return shape


def add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def write_para(tf, text, *, size, font=FONT_BODY, bold=False, color=INK, align=PP_ALIGN.LEFT, space_after=6, space_before=0, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, bold=bold, color=color)
    return p


def write_mixed(tf, parts, *, size, font=FONT_BODY, color=INK, align=PP_ALIGN.LEFT, space_after=8, first=False):
    """parts: list of (text, bold)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    for text, bold in parts:
        run = p.add_run()
        run.text = text
        set_run_font(run, font, size, bold=bold, color=color)
    return p


def add_notes(slide, text: str):
    notes = slide.notes_slide
    notes.notes_text_frame.text = text


def picture_size(path: Path, box_w, box_h):
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    box_ar = box_w / box_h
    if ar > box_ar:
        w = box_w
        h = int(box_w / ar)
    else:
        h = box_h
        w = int(box_h * ar)
    return w, h


def add_logo(slide, left, top, height):
    with Image.open(LOGO) as im:
        iw, ih = im.size
    width = int(height * (iw / ih))
    return slide.shapes.add_picture(str(LOGO), left, top, width=width, height=height)


def add_framed_shot(slide, path: Path, left, top, width, max_height, *, pad=Inches(0.08)):
    """Fit the bitmap at its native aspect, then vertically center in max_height."""
    inner_w = width - pad * 2
    pw, ph = picture_size(path, inner_w, max_height - pad * 2)
    height = ph + pad * 2
    if height > max_height:
        height = max_height
        pw, ph = picture_size(path, inner_w, max_height - pad * 2)
    top = top + (max_height - height) // 2
    add_rect(slide, left, top, width, height, TINT, rounded=True)
    px = left + pad + (inner_w - pw) // 2
    py = top + pad
    slide.shapes.add_picture(str(path), px, py, width=pw, height=ph)


def add_inner_chrome(slide, number: int, total: int):
    set_slide_bg(slide, PAPER)
    logo = add_logo(slide, MARGIN_X, Inches(0.22), Inches(0.42))
    kicker = add_textbox(slide, logo.left + logo.width + Inches(0.22), Inches(0.28), Inches(7.5), Inches(0.32))
    tf = kicker.text_frame
    tf.word_wrap = False
    write_para(
        tf,
        "Pengenalan  ·  Inspeksi DT",
        size=11,
        font=FONT_HEAD,
        bold=True,
        color=MUTED,
        first=True,
        space_after=0,
    )
    add_rect(slide, MARGIN_X, Inches(0.74), SLIDE_W - MARGIN_X * 2, Pt(2), ACCENT)
    foot_l = add_textbox(slide, MARGIN_X, FOOTER_Y, Inches(9.5), Inches(0.28))
    tf = foot_l.text_frame
    write_para(
        tf,
        "Inline Technology International",
        size=10,
        font=FONT_BODY,
        color=MUTED,
        first=True,
        space_after=0,
    )
    foot_r = add_textbox(slide, Inches(11.2), FOOTER_Y, Inches(1.55), Inches(0.28))
    tf = foot_r.text_frame
    write_para(
        tf,
        f"{number}  /  {total}",
        size=10,
        font=FONT_BODY,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        first=True,
        space_after=0,
    )


def style_table(table, rows, col_widths):
    table.first_row = False
    table.horz_banding = False
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.14)
            cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.08)
            cell.margin_bottom = Inches(0.08)
            fill = INK if r == 0 else (TINT if r % 2 == 0 else PAPER)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = text
            color = PAPER if r == 0 else INK
            set_run_font(
                run,
                FONT_HEAD if (r == 0 or c == 0) else FONT_BODY,
                13 if r == 0 else 14,
                bold=(r == 0 or c == 0),
                color=color,
            )


def add_title(slide, text: str, *, top=Inches(0.92), size=26):
    box = add_textbox(slide, MARGIN_X, top, Inches(12.2), Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    write_para(tf, text, size=size, font=FONT_HEAD, bold=True, color=INK, first=True, space_after=0)
    return box


def add_bullets(slide, left, top, width, height, items, *, size=16, space_after=10):
    box = add_textbox(slide, left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, list):
            write_mixed(tf, item, size=size, first=(i == 0), space_after=space_after)
        else:
            write_para(tf, item, size=size, first=(i == 0), space_after=space_after)
    return box


def add_card(slide, left, top, width, height, kicker: str, body: str):
    add_rect(slide, left, top, width, height, TINT, rounded=True)
    k = add_textbox(slide, left + Inches(0.18), top + Inches(0.14), width - Inches(0.32), Inches(0.32))
    tf = k.text_frame
    write_para(tf, kicker, size=11, font=FONT_HEAD, bold=True, color=INK, first=True, space_after=0)
    b = add_textbox(slide, left + Inches(0.18), top + Inches(0.46), width - Inches(0.32), height - Inches(0.58))
    tf = b.text_frame
    tf.word_wrap = True
    write_para(tf, body, size=13, color=INK, first=True, space_after=0)


def new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide_cover(prs: Presentation):
    slide = blank(prs)
    set_slide_bg(slide, INK)

    # 3×3 accent dots (shapes, not a rebuilt mark) — CSS nth-child(3n+2) uses tint
    grid_left = Inches(9.35)
    grid_top = Inches(1.55)
    cell = Inches(1.12)
    gap = Inches(0.22)
    size = Inches(0.9)
    for r in range(3):
        for c in range(3):
            idx = r * 3 + c  # 0-based
            fill = TINT if (idx % 3 == 1) else ACCENT
            add_oval(slide, grid_left + c * (cell + gap), grid_top + r * (cell + gap), size, fill)

    plate = add_rect(slide, MARGIN_X, Inches(0.42), Inches(4.35), Inches(1.35), PAPER, rounded=True)
    add_logo(slide, plate.left + Inches(0.28), plate.top + Inches(0.28), Inches(0.78))

    kicker = add_textbox(slide, MARGIN_X, Inches(2.15), Inches(8.4), Inches(0.36))
    tf = kicker.text_frame
    write_para(
        tf,
        "PANDUAN PENGENALAN",
        size=12,
        font=FONT_HEAD,
        bold=True,
        color=TINT,
        first=True,
        space_after=0,
    )

    title = add_textbox(slide, MARGIN_X, Inches(2.52), Inches(8.6), Inches(1.15))
    tf = title.text_frame
    tf.word_wrap = True
    write_para(tf, "Inspeksi DT", size=44, font=FONT_HEAD, bold=True, color=PAPER, first=True, space_after=0)

    product = add_textbox(slide, MARGIN_X, Inches(3.7), Inches(8.6), Inches(0.42))
    tf = product.text_frame
    write_para(
        tf,
        "Checklist walk-around digital untuk dump truck",
        size=18,
        font=FONT_HEAD,
        bold=True,
        color=ACCENT,
        first=True,
        space_after=0,
    )

    add_rect(slide, MARGIN_X, Inches(4.28), Inches(1.65), Pt(2.5), ACCENT)

    meta = add_textbox(slide, MARGIN_X, Inches(4.55), Inches(8.4), Inches(1.4))
    tf = meta.text_frame
    tf.word_wrap = True
    write_para(tf, "Inline Technology International", size=16, font=FONT_BODY, color=TINT, first=True, space_after=4)
    write_para(tf, "Versi aplikasi 1.8.0  ·  16 Sep 2026", size=14, color=TINT, space_after=4)
    write_para(tf, "Untuk semua peran  ·  ±10–15 menit", size=14, color=TINT, space_after=0)

    add_notes(slide, "Buka dengan identitas ITI; sebutkan bahwa ini pengenalan, bukan pelatihan teknis penuh.")
    return slide


def slide_agenda(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Agenda")
    items = [
        ("01", "Apa itu Inspeksi DT", "Checklist walk-around, foto bukti, persetujuan, rekap site."),
        ("02", "Siapa yang memakai", "Lima peran: Superadmin sampai Driver dan Mekanik."),
        ("03", "Alur Driver / Mekanik", "Unit dan tipe, checklist, foto, draf, kirim."),
        ("04", "Alur admin", "Persetujuan site, master perusahaan, rekap."),
        ("05", "Foto bukti", "Kamera in-app, cap GPS, batas 500 KB, masa simpan 90 hari."),
    ]
    col_w = Inches(5.85)
    for i, (num, title, body) in enumerate(items):
        col = i % 2
        row = i // 2
        left = MARGIN_X + col * (col_w + Inches(0.28))
        top = Inches(1.62) + row * Inches(1.62)
        add_rect(slide, left, top, col_w, Inches(1.42), TINT, rounded=True)
        add_rect(slide, left, top, Inches(0.08), Inches(1.42), ACCENT)
        num_box = add_textbox(slide, left + Inches(0.22), top + Inches(0.22), Inches(0.7), Inches(0.4))
        tf = num_box.text_frame
        write_para(tf, num, size=18, font=FONT_HEAD, bold=True, color=ACCENT, first=True, space_after=0)
        t = add_textbox(slide, left + Inches(0.95), top + Inches(0.22), col_w - Inches(1.2), Inches(0.38))
        tf = t.text_frame
        write_para(tf, title, size=16, font=FONT_HEAD, bold=True, first=True, space_after=0)
        b = add_textbox(slide, left + Inches(0.95), top + Inches(0.62), col_w - Inches(1.2), Inches(0.62))
        tf = b.text_frame
        tf.word_wrap = True
        write_para(tf, body, size=13, color=MUTED, first=True, space_after=0)
    add_notes(slide, "Janji waktu 10–15 menit; foto dan persetujuan dibahas singkat, bukan klik demi klik.")
    return slide


def slide_apa_itu(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Apa itu Inspeksi DT")
    cards = [
        ("Walk-around", "Checklist sesuai kategori kendaraan, dikurangi item yang dikecualikan pada tipe inspeksi."),
        ("Foto bukti", "Jepret di aplikasi. Cap waktu, GPS, dan nama pemeriksa tertanam di berkas."),
        ("Persetujuan", "Admin Site meninjau kiriman: Setujui atau Tolak. Setelah kirim, driver tidak mengedit."),
        ("Rekap site", "Matriks hijau / merah / abu per unit. Ekspor CSV untuk laporan site."),
    ]
    for i, (k, b) in enumerate(cards):
        left = MARGIN_X + (i % 2) * Inches(6.15)
        top = Inches(1.65) + (i // 2) * Inches(2.4)
        add_card(slide, left, top, Inches(5.95), Inches(2.18), k, b)
    add_notes(slide, "Tekankan bukti foto + jejak persetujuan, bukan sekadar centang kertas.")
    return slide


def slide_peran(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Siapa yang memakai")
    rows = [
        ("Peran", "Fokus", "Tidak untuk"),
        ("Superadmin", "Semua menu, termasuk Perusahaan", "—"),
        ("Admin Perusahaan", "Master checklist, tipe, site, rekap", "Akun Superadmin lain"),
        ("Admin Site", "Pengguna site, kendaraan, Setujui/Tolak, rekap", "Master checklist perusahaan"),
        ("Driver", "Inspeksi Baru, Inspeksi Saya", "Data master, rekap, persetujuan"),
        ("Mekanik", "Hak sama dengan Driver", "Sama dengan Driver"),
    ]
    table_shape = slide.shapes.add_table(len(rows), 3, MARGIN_X, Inches(1.58), Inches(12.22), Inches(4.85))
    style_table(table_shape.table, rows, [Inches(2.7), Inches(5.4), Inches(4.12)])
    add_notes(slide, "Mekanik = Driver. Admin Site tidak mengubah kategori/item/tipe inspeksi.")
    return slide


def slide_masuk(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Masuk ke aplikasi")
    add_bullets(
        slide,
        MARGIN_X,
        Inches(1.58),
        Inches(6.0),
        Inches(5.1),
        [
            [("Pilih bahasa ", False), ("EN / ID", True), (" di layar masuk.", False)],
            [("Akun dikunci setelah ", False), ("5 percobaan gagal", True), (" — tunggu ", False), ("15 menit", True), (".", False)],
            [("Tidak ada tautan ", False), ("lupa kata sandi", True), (". Minta admin mengubahnya di Pengguna & Peran.", False)],
            [("Minimum kata sandi ", False), ("6 karakter", True), (". Kosongkan field saat ubah pengguna jika tidak diganti.", False)],
            "Setelah masuk, menu mengikuti peran akun.",
        ],
        size=15,
        space_after=14,
    )
    add_framed_shot(slide, IMAGES / "01-login.png", Inches(6.85), Inches(1.55), Inches(5.9), Inches(5.15))
    add_notes(slide, "Jangan bagikan kata sandi produksi di sesi ini; demo hanya di slide terakhir.")
    return slide


def slide_driver_mulai(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Driver: mulai inspeksi")
    add_bullets(
        slide,
        MARGIN_X,
        Inches(1.58),
        Inches(6.05),
        Inches(5.1),
        [
            [("Pilih ", False), ("unit dan tipe inspeksi dulu", True), (". Checklist baru dimuat setelah keduanya terisi.", False)],
            [("Isi ", False), ("KM / HM", True), (" sesuai panel kendaraan.", False)],
            [("Ganti site, unit, atau tipe saat form sudah diisi: konfirmasi dulu — jawaban dibuang, KM tetap.", False)],
            "Driver/Mekanik, VIN, dan konfigurasi penggerak terisi otomatis.",
        ],
        size=15,
        space_after=14,
    )
    add_framed_shot(slide, IMAGES / "03-inspeksi-baru.png", Inches(6.85), Inches(1.55), Inches(5.9), Inches(5.15))
    add_notes(slide, "Tekankan gerbang unit + tipe. Jangan mulai isi checklist sebelum keduanya dipilih.")
    return slide


def slide_driver_checklist(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Driver: isi checklist")
    add_bullets(
        slide,
        MARGIN_X,
        Inches(1.58),
        Inches(6.05),
        Inches(5.1),
        [
            [("Setiap item dimulai ", False), ("OK", True), (". Ubah hanya yang bermasalah.", False)],
            [("Catatan dan foto ", False), ("opsional", True), (" — termasuk pada item yang tetap OK.", False)],
            [("Lampu kilat (torch) ", False), ("jika kamera mendukung", True), (". Banyak komputer desktop tidak punya tombol ini.", False)],
            "Unggah berkas dari galeri tidak memakai lampu kilat.",
            "Kategori yang seluruh itemnya dikecualikan pada tipe tidak muncul.",
        ],
        size=15,
        space_after=12,
    )
    add_framed_shot(slide, IMAGES / "04-checklist.png", Inches(6.85), Inches(1.55), Inches(5.9), Inches(5.15))
    add_notes(slide, "Default OK mempercepat walk-around; foto tetap boleh pada item OK.")
    return slide


def slide_driver_foto(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Driver: foto bukti")
    cards = [
        ("Jepret", "Buka kamera in-app. Izinkan akses kamera di peramban saat diminta."),
        ("Ulangi", "Ambil ulang jika buram atau salah sudut. Pratinjau tetap hidup."),
        ("Gunakan foto", "Kompresi otomatis, maksimum 500 KB, lalu tertanam di item."),
        ("Cap bukti", "Tanggal/waktu, GPS atau “GPS tidak tersedia”, lokasi, nama pemeriksa."),
    ]
    for i, (k, b) in enumerate(cards):
        left = MARGIN_X + (i % 2) * Inches(6.15)
        top = Inches(1.62) + (i // 2) * Inches(2.42)
        add_card(slide, left, top, Inches(5.95), Inches(2.2), k, b)
    add_notes(slide, "Torch = lampu terus-menerus, bukan kilat sekali. GPS boleh “tidak tersedia”.")
    return slide


def slide_driver_draf(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Driver: draf dan kirim")
    add_bullets(
        slide,
        MARGIN_X,
        Inches(1.58),
        Inches(6.05),
        Inches(5.1),
        [
            [("Simpan progres", True), (" menyimpan draf di perangkat ini saja — bukan di server.", False)],
            [("Daftar draf: ", False), ("Inspeksi Baru", True), (" dan ", False), ("Inspeksi Saya", True), (". Lanjutkan atau Hapus draf.", False)],
            "Draf hilang jika ganti peramban, hapus data situs, atau pindah ponsel.",
            [("Setelah ", False), ("Kirim", True), (", inspeksi tidak bisa diedit. Status: Terkirim → Disetujui / Ditolak.", False)],
        ],
        size=15,
        space_after=14,
    )
    add_framed_shot(slide, IMAGES / "02-beranda.png", Inches(6.85), Inches(1.55), Inches(5.9), Inches(5.15))
    add_notes(slide, "Inspeksi Saya = kiriman server + draf lokal. Jangan andalkan draf sebagai arsip perusahaan.")
    return slide


def slide_admin_site(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Admin Site: setujui atau tolak")
    add_bullets(
        slide,
        MARGIN_X,
        Inches(1.58),
        Inches(6.05),
        Inches(5.1),
        [
            [("Buka ", False), ("Laporan Inspeksi", True), (", pilih kiriman, tinjau checklist dan foto.", False)],
            [("Setujui", True), (" atau ", False), ("Tolak", True), (". Keputusan ini menutup siklus driver.", False)],
            "Kelola akun Driver dan Mekanik untuk site sendiri — bukan Admin Perusahaan.",
            "Kelola daftar kendaraan site. Master checklist (kategori/item/tipe) milik perusahaan.",
        ],
        size=15,
        space_after=14,
    )
    add_framed_shot(slide, IMAGES / "06-detail-persetujuan.png", Inches(6.85), Inches(1.55), Inches(5.9), Inches(5.15))
    add_notes(slide, "Admin Site = operasional harian. Jangan janji mereka bisa mengubah item checklist.")
    return slide


def slide_admin_perusahaan(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Admin Perusahaan: checklist dan tipe")
    add_bullets(
        slide,
        MARGIN_X,
        Inches(1.58),
        Inches(6.05),
        Inches(5.1),
        [
            [("Urutan master: ", False), ("item → kategori → tipe", True), (". Kategori kendaraan memasang kumpulan item.", False)],
            [("Kecualikan item", True), (" pada tipe: dikurangi dari kumpulan itu. Daftar kosong = checklist penuh.", False)],
            "Kategori yang seluruh itemnya dikecualikan tidak tampil di form.",
            "Buat Admin Site, Driver, dan Mekanik. Tidak membuat Superadmin atau Admin Perusahaan lain.",
        ],
        size=14,
        space_after=12,
    )
    add_framed_shot(slide, IMAGES / "11-tipe-kecualikan.png", Inches(6.85), Inches(1.55), Inches(5.9), Inches(5.15))
    add_notes(slide, "Tipe adalah daftar pengecualian, bukan daftar item mandiri.")
    return slide


def slide_rekap(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Rekap kendaraan site")
    add_bullets(
        slide,
        MARGIN_X,
        Inches(1.58),
        Inches(6.05),
        Inches(5.1),
        [
            [("Matriks Harian", True), (": kisi unit × tanggal. Hijau = Diinspeksi, semua OK; merah = Ada temuan; abu = Belum diinspeksi.", False)],
            "Klik sel untuk membuka detail. Beberapa inspeksi di hari yang sama: sel mengikuti inspeksi terakhir; merah jika ada temuan.",
            [("Ekspor CSV Matriks", True), (" dan ", False), ("Ekspor CSV Ringkasan", True), (". Laporan Inspeksi punya Ekspor CSV mengikuti filter.", False)],
            "Admin perusahaan memilih Site di dalam perusahaannya. Superadmin wajib pilih site.",
        ],
        size=15,
        space_after=12,
    )
    add_framed_shot(slide, IMAGES / "07-rekap.png", Inches(6.85), Inches(1.55), Inches(5.9), Inches(5.15))
    add_notes(slide, "Jangan overclaim arti warna di luar SOP manual; CSV untuk operasional site.")
    return slide


def slide_ingat(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Yang perlu diingat")
    points = [
        ("Checklist = kategori kendaraan minus exclude tipe", "Form hanya memuat item kategori unit, dikurangi Kecualikan item pada tipe yang dipilih."),
        ("Draf hanya di perangkat", "Simpan progres tidak naik ke server. Ganti HP atau hapus data situs = draf hilang."),
        ("Foto 90 hari", "Masa simpan default 90 hari. Data inspeksi tetap ada; berkas foto bisa dibersihkan."),
        ("Setelah kirim, selesai", "Tidak ada sunting kiriman. Koreksi lewat inspeksi baru atau keputusan admin."),
    ]
    for i, (k, b) in enumerate(points):
        top = Inches(1.55) + i * Inches(1.25)
        add_rect(slide, MARGIN_X, top, Inches(0.12), Inches(1.08), ACCENT, rounded=True)
        add_rect(slide, MARGIN_X + Inches(0.12), top, Inches(12.1), Inches(1.08), TINT)
        t = add_textbox(slide, MARGIN_X + Inches(0.38), top + Inches(0.12), Inches(11.6), Inches(0.36))
        tf = t.text_frame
        write_para(tf, k, size=16, font=FONT_HEAD, bold=True, first=True, space_after=0)
        d = add_textbox(slide, MARGIN_X + Inches(0.38), top + Inches(0.5), Inches(11.6), Inches(0.48))
        tf = d.text_frame
        tf.word_wrap = True
        write_para(tf, b, size=13, color=MUTED, first=True, space_after=0)
    add_notes(slide, "Tiga pesan yang harus tertinggal: exclude tipe, draf lokal, foto 90 hari.")
    return slide


def slide_latihan(prs, n, total):
    slide = blank(prs)
    add_inner_chrome(slide, n, total)
    add_title(slide, "Latihan lokal saja")

    banner = add_rect(slide, MARGIN_X, Inches(1.55), Inches(12.22), Inches(0.72), TINT, rounded=True)
    _ = banner
    bt = add_textbox(slide, MARGIN_X + Inches(0.28), Inches(1.68), Inches(11.7), Inches(0.48))
    tf = bt.text_frame
    tf.word_wrap = True
    write_para(
        tf,
        "Bukan SOP produksi. Akun di bawah hanya untuk lingkungan demo/latihan. Jangan dipakai di server live.",
        size=14,
        font=FONT_HEAD,
        bold=True,
        color=WARN,
        first=True,
        space_after=0,
    )

    rows = [
        ("Peran", "Email", "Kata sandi (demo)"),
        ("Driver", "driver@iti.demo", "Driver@1234"),
        ("Admin Site", "admin@iti.demo", "Admin@1234"),
        ("Admin Perusahaan", "company.admin@iti.demo", "Admin@1234"),
    ]
    table_shape = slide.shapes.add_table(len(rows), 3, MARGIN_X, Inches(2.5), Inches(12.22), Inches(3.15))
    style_table(table_shape.table, rows, [Inches(3.2), Inches(5.4), Inches(3.62)])

    note = add_textbox(slide, MARGIN_X, Inches(5.85), Inches(12.22), Inches(1.0))
    tf = note.text_frame
    tf.word_wrap = True
    write_para(
        tf,
        "Superadmin tidak memakai panel akun demo (dibuat dari konfigurasi server). Mekanik seed mechanic@iti.demo tidak tampil di panel Akun demo. Frontend latihan: localhost:3000.",
        size=13,
        color=MUTED,
        first=True,
        space_after=0,
    )
    add_notes(slide, "Tegaskan slide ini latihan lokal. Tutup tanpa meninggalkan kata sandi di papan jika sesi hybrid.")
    return slide


def build():
    if not LOGO.exists():
        raise SystemExit(f"Logo not found: {LOGO}")
    prs = new_prs()
    builders = [
        lambda: slide_cover(prs),
        lambda: slide_agenda(prs, 2, 14),
        lambda: slide_apa_itu(prs, 3, 14),
        lambda: slide_peran(prs, 4, 14),
        lambda: slide_masuk(prs, 5, 14),
        lambda: slide_driver_mulai(prs, 6, 14),
        lambda: slide_driver_checklist(prs, 7, 14),
        lambda: slide_driver_foto(prs, 8, 14),
        lambda: slide_driver_draf(prs, 9, 14),
        lambda: slide_admin_site(prs, 10, 14),
        lambda: slide_admin_perusahaan(prs, 11, 14),
        lambda: slide_rekap(prs, 12, 14),
        lambda: slide_ingat(prs, 13, 14),
        lambda: slide_latihan(prs, 14, 14),
    ]
    for fn in builders:
        fn()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
