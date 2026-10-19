#!/usr/bin/env python3
"""Customer adoption deck (16:9). Does not replace the how-to deck.

Bahasa Indonesia titles, one English subtitle each. Screenshots are the
ones already saved under images/intro/.

Rebuild:
    python3 docs/panduan-pengguna/build_sales_pptx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

import build_intro_pptx as ui  # noqa: E402

from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

OUT = ROOT.parent / "Inspeksi_Aset_Adopsi.pptx"
KICKER = "Adopsi  ·  Inspeksi Aset"
TOTAL = 11


def chrome(slide, number: int):
    ui.add_inner_chrome(slide, number, TOTAL, KICKER)


def pair(slide, left_path, right_path, left_cap: str, right_cap: str):
    top = Inches(2.05)
    width = Inches(6.0)
    height = Inches(4.55)
    gap = Inches(0.22)
    ui.add_framed_shot(slide, left_path, ui.MARGIN_X, top, width, height)
    ui.add_framed_shot(slide, right_path, ui.MARGIN_X + width + gap, top, width, height)
    for i, cap in enumerate((left_cap, right_cap)):
        box = ui.add_textbox(slide, ui.MARGIN_X + i * (width + gap), Inches(1.72), width, Inches(0.3))
        tf = box.text_frame
        ui.write_para(tf, cap, size=12, font=ui.FONT_HEAD, bold=True, color=ui.MUTED, first=True, space_after=0)


def slide_cover(prs):
    slide = ui.blank(prs)
    ui.set_slide_bg(slide, ui.INK)

    grid_left = Inches(9.35)
    grid_top = Inches(1.55)
    cell = Inches(1.12)
    gap = Inches(0.22)
    size = Inches(0.9)
    for r in range(3):
        for c in range(3):
            idx = r * 3 + c
            fill = ui.TINT if (idx % 3 == 1) else ui.ACCENT
            ui.add_oval(slide, grid_left + c * (cell + gap), grid_top + r * (cell + gap), size, fill)

    plate = ui.add_rect(slide, ui.MARGIN_X, Inches(0.42), Inches(4.35), Inches(1.35), ui.PAPER, rounded=True)
    ui.add_logo(slide, plate.left + Inches(0.28), plate.top + Inches(0.28), Inches(0.78))

    kicker = ui.add_textbox(slide, ui.MARGIN_X, Inches(2.05), Inches(8.6), Inches(0.36))
    tf = kicker.text_frame
    ui.write_para(tf, "MENGAPA BERGANTI", size=12, font=ui.FONT_HEAD, bold=True, color=ui.TINT, first=True, space_after=0)

    title = ui.add_textbox(slide, ui.MARGIN_X, Inches(2.45), Inches(8.8), Inches(1.05))
    tf = title.text_frame
    tf.word_wrap = True
    ui.write_para(tf, "Inspeksi Aset", size=44, font=ui.FONT_HEAD, bold=True, color=ui.PAPER, first=True, space_after=0)

    product = ui.add_textbox(slide, ui.MARGIN_X, Inches(3.55), Inches(8.6), Inches(1.15))
    tf = product.text_frame
    tf.word_wrap = True
    ui.write_para(tf, "Asset Inspection", size=20, font=ui.FONT_HEAD, bold=True, color=ui.ACCENT, first=True, space_after=4)
    ui.write_para(
        tf,
        "Ganti keliling kertas dengan catatan yang bisa dipercaya kantor.",
        size=16,
        color=ui.TINT,
        space_after=0,
    )

    ui.add_rect(slide, ui.MARGIN_X, Inches(4.9), Inches(1.65), Pt(2.5), ui.ACCENT)
    meta = ui.add_textbox(slide, ui.MARGIN_X, Inches(5.15), Inches(8.4), Inches(1.1))
    tf = meta.text_frame
    tf.word_wrap = True
    ui.write_para(tf, "Inline Technology International", size=16, font=ui.FONT_BODY, color=ui.TINT, first=True, space_after=4)
    ui.write_para(tf, "15 Oktober 2026  ·  untuk pelanggan", size=14, color=ui.TINT, space_after=0)
    ui.add_notes(slide, "Jangan mulai dari cara klik. Mulai dari apa yang hilang hari ini.")
    return slide


def slide_gap(prs):
    slide = ui.blank(prs)
    chrome(slide, 2)
    ui.add_title(slide, "Yang hilang hari ini", "Paper, chat photos, and a recap built by hand")
    cards = [
        ("Kertas", "Checklist mudah tercecer. Besok tidak jelas unit mana yang benar-benar dikelilingi."),
        ("Foto di chat", "Gambar lepas dari item. Sulit mencari bukti saat ada temuan."),
        ("Tanpa keputusan", "Tidak ada jejak siapa yang menyetujui, atau kapan."),
        ("Rekap manual", "Angka site disusun dari banyak sumber, sering terlambat."),
    ]
    for i, (k, b) in enumerate(cards):
        left = ui.MARGIN_X + (i % 2) * Inches(6.15)
        top = Inches(1.78) + (i // 2) * Inches(2.35)
        ui.add_card(slide, left, top, Inches(5.95), Inches(2.15), k, b)
    return slide


def slide_changes(prs):
    slide = ui.blank(prs)
    chrome(slide, 3)
    ui.add_title(slide, "Yang berubah", "One inspection, one photo, one decision, one matrix")
    steps = [
        ("01", "Satu inspeksi", "Unit, tipe, dan KM/HM tersimpan sebagai satu catatan."),
        ("02", "Foto di item", "Bukti menempel pada baris yang diperiksa, bukan di percakapan."),
        ("03", "Satu keputusan", "Admin site menyetujui atau menolak kiriman yang sama."),
        ("04", "Matriks hari itu", "Site melihat unit mana yang hijau, merah, atau belum dikelilingi."),
    ]
    for i, (num, title, body) in enumerate(steps):
        left = ui.MARGIN_X + (i % 2) * Inches(6.15)
        top = Inches(1.78) + (i // 2) * Inches(2.35)
        ui.add_rect(slide, left, top, Inches(5.95), Inches(2.15), ui.TINT, rounded=True)
        ui.add_rect(slide, left, top, Inches(0.1), Inches(2.15), ui.ACCENT)
        n = ui.add_textbox(slide, left + Inches(0.28), top + Inches(0.28), Inches(1.1), Inches(0.4))
        tf = n.text_frame
        ui.write_para(tf, num, size=18, font=ui.FONT_HEAD, bold=True, color=ui.ACCENT, first=True, space_after=0)
        t = ui.add_textbox(slide, left + Inches(1.4), top + Inches(0.28), Inches(4.2), Inches(0.4))
        tf = t.text_frame
        ui.write_para(tf, title, size=18, font=ui.FONT_HEAD, bold=True, first=True, space_after=0)
        b = ui.add_textbox(slide, left + Inches(0.28), top + Inches(0.9), Inches(5.35), Inches(0.95))
        tf = b.text_frame
        tf.word_wrap = True
        ui.write_para(tf, body, size=15, color=ui.INK, first=True, space_after=0)
    return slide


def slide_walk(prs):
    slide = ui.blank(prs)
    chrome(slide, 4)
    ui.add_title(slide, "Keliling unit, sekali jalan", "The yard fills one form, not a stack of paper")
    pair(
        slide,
        ui.INTRO / "form.png",
        ui.INTRO / "checklist.png",
        "Unit, tipe, dan KM/HM",
        "Setiap item mulai OK",
    )
    return slide


def slide_evidence(prs):
    slide = ui.blank(prs)
    chrome(slide, 5)
    ui.add_title(slide, "Bukti yang bisa dibuka lagi", "The finding, the note, and the photo stay together")
    ui.add_bullets(
        slide,
        ui.MARGIN_X,
        Inches(1.78),
        Inches(6.05),
        Inches(4.9),
        [
            "Status masih Terkirim sampai ada keputusan.",
            "Catatan temuan tertulis di item, bukan hanya di chat.",
            "Foto menempel di baris itu, dengan waktu dan nama pemeriksa.",
            "Kantor membuka catatan yang sama dengan yang diisi di unit.",
        ],
        size=15,
        space_after=12,
    )
    ui.add_framed_shot(slide, ui.INTRO / "detail.png", Inches(6.85), Inches(1.72), Inches(5.9), Inches(5.05))
    return slide


def slide_decision(prs):
    slide = ui.blank(prs)
    chrome(slide, 6)
    ui.add_title(slide, "Keputusan di hari yang sama", "Approve or reject, then keep a PDF")
    cards = [
        ("Setujui", "Admin site menutup siklus bila checklist dan foto sudah cukup."),
        ("Tolak", "Kiriman yang belum layak tidak tercampur dengan yang sudah beres."),
        ("PDF", "Satu berkas A4 untuk arsip: ringkasan, item, dan foto temuan."),
        ("Tetap terbuka", "Tombol keputusan hanya ada selagi status masih Terkirim."),
    ]
    for i, (k, b) in enumerate(cards):
        left = ui.MARGIN_X + (i % 2) * Inches(6.15)
        top = Inches(1.78) + (i // 2) * Inches(2.35)
        ui.add_card(slide, left, top, Inches(5.95), Inches(2.15), k, b)
    return slide


def slide_site(prs):
    slide = ui.blank(prs)
    chrome(slide, 7)
    ui.add_title(slide, "Site tanpa mengejar driver", "Pending work and the daily matrix, without a chase")
    pair(
        slide,
        ui.INTRO / "dashboard.png",
        ui.INTRO / "recap.png",
        "Berapa yang masih menunggu",
        "Hijau, merah, atau belum dikelilingi",
    )
    return slide


def slide_yard(prs):
    slide = ui.blank(prs)
    chrome(slide, 8)
    ui.add_title(slide, "Dibuat untuk lapangan", "The yard does not stop when the signal drops")
    points = [
        ("Draf di perangkat", "Progres tersimpan di ponsel ini. Belum naik ke server sampai dikirim atau disinkronkan."),
        ("Antrean offline", "Driver, mekanik, dan admin site mengantre inspeksi, lalu menekan Sinkron saat tersambung."),
        ("Aplikasi Android", "Alur yang sama di ponsel. PDF dibagikan lewat lembar berbagi Android."),
        ("Kantor tetap online", "Admin perusahaan dan superadmin bekerja saat tersambung, di atas data yang sudah masuk."),
    ]
    for i, (k, b) in enumerate(points):
        top = Inches(1.72) + i * Inches(1.22)
        ui.add_rect(slide, ui.MARGIN_X, top, Inches(0.12), Inches(1.06), ui.ACCENT, rounded=True)
        ui.add_rect(slide, ui.MARGIN_X + Inches(0.12), top, Inches(12.1), Inches(1.06), ui.TINT)
        t = ui.add_textbox(slide, ui.MARGIN_X + Inches(0.38), top + Inches(0.1), Inches(11.6), Inches(0.34))
        tf = t.text_frame
        ui.write_para(tf, k, size=16, font=ui.FONT_HEAD, bold=True, first=True, space_after=0)
        d = ui.add_textbox(slide, ui.MARGIN_X + Inches(0.38), top + Inches(0.46), Inches(11.6), Inches(0.5))
        tf = d.text_frame
        tf.word_wrap = True
        ui.write_para(tf, b, size=14, color=ui.MUTED, first=True, space_after=0)
    return slide


def slide_roles(prs):
    slide = ui.blank(prs)
    chrome(slide, 9)
    ui.add_title(slide, "Siapa memegang apa", "Four jobs, not a training manual")
    rows = [
        ("Driver / Mekanik", "Mengisi, menyimpan, dan mengirim inspeksi unit."),
        ("Admin site", "Melihat seluruh site, lalu menyetujui atau menolak."),
        ("Admin perusahaan", "Menjaga checklist, tipe, dan rekap di perusahaannya."),
        ("Superadmin", "Membuka semua perusahaan bila organisasi punya lebih dari satu."),
    ]
    for i, (role, body) in enumerate(rows):
        top = Inches(1.78) + i * Inches(1.18)
        ui.add_rect(slide, ui.MARGIN_X, top, Inches(12.22), Inches(1.05), ui.TINT, rounded=True)
        r = ui.add_textbox(slide, ui.MARGIN_X + Inches(0.28), top + Inches(0.28), Inches(3.3), Inches(0.5))
        tf = r.text_frame
        ui.write_para(tf, role, size=16, font=ui.FONT_HEAD, bold=True, first=True, space_after=0)
        b = ui.add_textbox(slide, ui.MARGIN_X + Inches(3.7), top + Inches(0.26), Inches(8.1), Inches(0.55))
        tf = b.text_frame
        tf.word_wrap = True
        ui.write_para(tf, body, size=15, color=ui.INK, first=True, space_after=0)
    return slide


def slide_start(prs):
    slide = ui.blank(prs)
    chrome(slide, 10)
    ui.add_title(slide, "Cara mulai", "One site and the daily check, then grow")
    ui.add_card(
        slide,
        ui.MARGIN_X,
        Inches(1.78),
        Inches(6.0),
        Inches(4.85),
        "Urutan",
        "",
    )
    steps = ui.add_textbox(slide, ui.MARGIN_X + Inches(0.22), Inches(2.35), Inches(5.55), Inches(3.9))
    tf = steps.text_frame
    tf.word_wrap = True
    lines = [
        "1. Satu site, tipe harian (P2H).",
        "2. Driver mengirim, admin site memutuskan.",
        "3. Rekap dipakai di rapat site.",
        "4. Tambah tipe dan site bila pola ini sudah jalan.",
    ]
    for i, line in enumerate(lines):
        ui.write_para(tf, line, size=16, first=(i == 0), space_after=14)

    ui.add_rect(slide, ui.MARGIN_X + Inches(6.25), Inches(1.78), Inches(5.95), Inches(2.3), ui.TINT, rounded=True)
    keep = ui.add_textbox(slide, ui.MARGIN_X + Inches(6.48), Inches(1.95), Inches(5.5), Inches(1.95))
    tf = keep.text_frame
    tf.word_wrap = True
    ui.write_para(tf, "Tetap di perusahaan", size=14, font=ui.FONT_HEAD, bold=True, first=True, space_after=6)
    ui.write_para(tf, "Master checklist, tipe inspeksi, pengguna, dan daftar unit.", size=14, color=ui.INK, space_after=0)

    ui.add_rect(slide, ui.MARGIN_X + Inches(6.25), Inches(4.25), Inches(5.95), Inches(2.38), ui.INK, rounded=True)
    stop = ui.add_textbox(slide, ui.MARGIN_X + Inches(6.48), Inches(4.42), Inches(5.5), Inches(2.0))
    tf = stop.text_frame
    tf.word_wrap = True
    ui.write_para(tf, "Berhenti di kertas", size=14, font=ui.FONT_HEAD, bold=True, color=ui.PAPER, first=True, space_after=6)
    ui.write_para(
        tf,
        "Keliling kertas, foto yang hanya ada di chat, dan rekap yang disusun tangan setiap sore.",
        size=14,
        color=ui.TINT,
        space_after=0,
    )
    return slide


def slide_close(prs):
    slide = ui.blank(prs)
    chrome(slide, 11)
    ui.add_title(slide, "Tiga bukti untuk dibawa pulang", "Leave with evidence, a decision, and the day’s matrix")
    points = [
        ("01", "Bukti di item", "Foto dan catatan tidak tercecer. Keduanya menempel pada baris yang diperiksa."),
        ("02", "Satu keputusan", "Setujui atau Tolak menutup kiriman. PDF menyimpan salinan untuk arsip."),
        ("03", "Rekap hari itu", "Site melihat unit hijau, merah, atau belum dikelilingi, tanpa mengejar driver."),
    ]
    for i, (num, title, body) in enumerate(points):
        top = Inches(1.85) + i * Inches(1.55)
        ui.add_rect(slide, ui.MARGIN_X, top, Inches(12.22), Inches(1.4), ui.TINT, rounded=True)
        n = ui.add_textbox(slide, ui.MARGIN_X + Inches(0.28), top + Inches(0.38), Inches(0.9), Inches(0.5))
        tf = n.text_frame
        ui.write_para(tf, num, size=22, font=ui.FONT_HEAD, bold=True, color=ui.ACCENT, align=PP_ALIGN.LEFT, first=True, space_after=0)
        t = ui.add_textbox(slide, ui.MARGIN_X + Inches(1.3), top + Inches(0.22), Inches(10.5), Inches(0.4))
        tf = t.text_frame
        ui.write_para(tf, title, size=18, font=ui.FONT_HEAD, bold=True, first=True, space_after=0)
        b = ui.add_textbox(slide, ui.MARGIN_X + Inches(1.3), top + Inches(0.68), Inches(10.5), Inches(0.55))
        tf = b.text_frame
        tf.word_wrap = True
        ui.write_para(tf, body, size=14, color=ui.MUTED, first=True, space_after=0)
    return slide


def build():
    missing = [name for name in ("dashboard.png", "form.png", "checklist.png", "detail.png", "recap.png") if not (ui.INTRO / name).exists()]
    if missing:
        raise SystemExit(f"Missing screenshots in {ui.INTRO}: {', '.join(missing)}")
    if not ui.LOGO.exists():
        raise SystemExit(f"Logo not found: {ui.LOGO}")
    prs = ui.new_prs()
    for fn in (
        slide_cover,
        slide_gap,
        slide_changes,
        slide_walk,
        slide_evidence,
        slide_decision,
        slide_site,
        slide_yard,
        slide_roles,
        slide_start,
        slide_close,
    ):
        fn(prs)
    prs.save(str(OUT))
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
